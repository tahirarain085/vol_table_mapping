"""Tree enumeration register (Flask + SQLAlchemy + MySQL on Aiven) with user accounts.

Run locally:   python app.py
Production:    gunicorn app:app
"""
import hmac
import os
import re
import secrets
import ssl
import tempfile
import time
import uuid
from collections import defaultdict
from datetime import datetime
from functools import wraps
from io import BytesIO
from urllib.parse import parse_qsl, quote_plus, urlencode, urlsplit, urlunsplit

from dotenv import load_dotenv
from flask import (
    Flask, abort, flash, g, redirect, render_template, render_template_string,
    request, send_file, session, url_for,
)
from flask_sqlalchemy import SQLAlchemy
from markupsafe import escape
from openpyxl import Workbook
from openpyxl.styles import Border, Font, Side
from sqlalchemy import func, inspect, text
from sqlalchemy.exc import IntegrityError, OperationalError
from werkzeug.exceptions import HTTPException
from werkzeug.middleware.proxy_fix import ProxyFix
from werkzeug.security import check_password_hash, generate_password_hash

from volume_data import RANGES

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
load_dotenv(os.path.join(BASE_DIR, ".env"))  # must run BEFORE any os.getenv


# ---------------------------------------------------------------------------
# Database configuration
#   1) DATABASE_URL  (paste the "Service URI" from Aiven), or
#   2) DB_USERNAME / DB_PASSWORD / DB_HOST / DB_PORT / DB_NAME
#   3) nothing set -> a local SQLite file, so the app still runs for testing
# ---------------------------------------------------------------------------
def _mysql_ssl_args():
    """SSL settings for Aiven. Uses ca.pem if present, else encrypts without verifying."""
    if os.getenv("DB_SSL", "true").strip().lower() in ("0", "false", "no", "off"):
        return {}
    ca_path = (os.getenv("DB_SSL_CA") or "").strip() or os.path.join(BASE_DIR, "ca.pem")
    pem_text = os.getenv("DB_SSL_CA_PEM")
    if pem_text and not os.path.exists(ca_path):
        ca_path = os.path.join(tempfile.gettempdir(), "aiven_ca.pem")
        with open(ca_path, "w") as f:
            f.write(pem_text.replace("\\n", "\n"))
    if os.path.exists(ca_path):
        return {"ssl": {"ca": ca_path}}
    print("[database] WARNING: ca.pem not found. Connection is encrypted but the server "
          "certificate is NOT verified. Download ca.pem from Aiven and put it next to app.py.")
    ctx = ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = ssl.CERT_NONE
    return {"ssl": ctx}


def build_database_config():
    """Returns (uri, engine_options)."""
    url = (os.getenv("DATABASE_URL") or "").strip()
    if url:
        url = re.sub(r"^mysql://", "mysql+pymysql://", url)
        parts = urlsplit(url)
        query = [(k, v) for k, v in parse_qsl(parts.query) if k.lower().replace("_", "-") != "ssl-mode"]
        if not any(k == "charset" for k, _ in query):
            query.append(("charset", "utf8mb4"))
        url = urlunsplit(parts._replace(query=urlencode(query)))
    else:
        host = (os.getenv("DB_HOST") or "").strip()
        if not host:
            print("[database] No DB_HOST or DATABASE_URL set: using local SQLite file trees_local.db")
            return f"sqlite:///{os.path.join(BASE_DIR, 'trees_local.db')}", {}
        user = (os.getenv("DB_USERNAME") or "").strip()
        password = quote_plus(os.getenv("DB_PASSWORD") or "")
        port = (os.getenv("DB_PORT") or "3306").strip()
        name = (os.getenv("DB_NAME") or "").strip()
        url = f"mysql+pymysql://{user}:{password}@{host}:{port}/{name}?charset=utf8mb4"

    options = {"pool_pre_ping": True, "pool_recycle": 280, "connect_args": _mysql_ssl_args()}
    return url, options


app = Flask(__name__)

SECRET_KEY = os.getenv("SECRET_KEY")
if not SECRET_KEY:
    raise RuntimeError("SECRET_KEY is not set. Add it to .env or to Render > Environment.")
app.config["SECRET_KEY"] = SECRET_KEY
app.config["SQLALCHEMY_DATABASE_URI"], app.config["SQLALCHEMY_ENGINE_OPTIONS"] = build_database_config()
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False

# Sign-in cookie: not readable by scripts, not sent from other sites, HTTPS only on Render
ON_RENDER = bool(os.getenv("RENDER"))
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = ON_RENDER or os.getenv("SESSION_COOKIE_SECURE", "0") == "1"
if ON_RENDER:  # Render sits in front of the app: trust it for the visitor's IP and https
    app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)

db = SQLAlchemy(app)


# ---------------------------------------------------------------------------
# Beat names offered in the dropdown. Every user can also type any other name.
# ---------------------------------------------------------------------------
BEATS = [
    "Rajana Kamalia Road",
    "Toba Rajana Road",
    "Moongi Beat",
    "Rajana Pirmahal Road",
    "Rajana Samundri Road",
    "Bhagat Canal",
    "Pirmahal Beat 1"
]


def build_volume_table(ranges):
    """Expand RANGES into {(kind_tree, girth): {...}} and check for mistakes."""
    table = {}
    for i, row in enumerate(ranges, start=1):
        if len(row) != 6:
            raise ValueError(f"RANGES row {i} must have 6 values, found {len(row)}: {row}")

        kind, g_from, g_to, dia, cft, shi_unit = row
        if g_from > g_to:
            raise ValueError(f"RANGES row {i}: girth_from is bigger than girth_to: {row}")

        for girth in range(g_from, g_to + 1):
            key = (kind.strip().lower(), girth)
            if key in table:
                raise ValueError(f"RANGES row {i}: girth {girth} of {kind} is defined twice")
            table[key] = {"dia": dia, "cft": cft, "shi_unit": shi_unit}
    return table


VOLUME_TABLE = build_volume_table(RANGES)
KINDS = sorted({kind for kind, _ in VOLUME_TABLE})


# ---------------------------------------------------------------------------
# Tree no rule
# The same tree no CAN be used again in another km / RD no (or another beat),
# but it can NOT be saved twice in the same km / RD no of the same beat.
# (Every beat has its own table, so the rule is checked inside that table.)
# Set TREE_NO_UNIQUE_PER_KIND = True if the same tree no may also be used once
# for each kind (for example shisham 14 and kikar 14) in the same km / RD no.
# ---------------------------------------------------------------------------
TREE_NO_UNIQUE_PER_KIND = False

# Location is optional (GPS can fail under thick canopy or without a signal).
# Set to True to refuse a tree that has no coordinates.
REQUIRE_LOCATION = False
UNIQUE_COLUMNS = ["km_burji_no"] + (["kind_tree"] if TREE_NO_UNIQUE_PER_KIND else []) + ["tree_no"]


# ---------------------------------------------------------------------------
# Tables
#   "app_user"                    : the accounts
#   "beats"                       : the beats (id, owner, name, table_name)
#   "beat_<id>_<name>"            : one table per beat, holds only that beat's trees
# Every beat belongs to ONE user, so every tree belongs to that user too.
# Table names are built from the beat's number and letters a-z / 0-9 only,
# never from raw text, so a beat name can not break the SQL.
# ---------------------------------------------------------------------------
class User(db.Model):
    __tablename__ = "app_user"

    id = db.Column(db.Integer, primary_key=True)
    full_name = db.Column(db.String(100), nullable=False)
    username = db.Column(db.String(50), nullable=False, unique=True)
    email = db.Column(db.String(120), unique=True)
    password_hash = db.Column(db.String(255), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.now)


class Beat(db.Model):
    __tablename__ = "beats"

    id = db.Column(db.Integer, primary_key=True)
    # Owner. Empty only for beats made before accounts existed (the first account takes them).
    user_id = db.Column(db.Integer, db.ForeignKey("app_user.id"), nullable=True)
    name = db.Column(db.String(100), nullable=False)
    table_name = db.Column(db.String(80), nullable=False, unique=True)

    # Two users may both have a beat with the same name, but one user may not have it twice
    __table_args__ = (db.UniqueConstraint("user_id", "name", name="uq_beats_user_name"),)


_tree_models = {}


def tree_model(table_name):
    """Return the model class for one beat's table (made once, then reused)."""
    model = _tree_models.get(table_name)
    if model is None:
        attrs = {
            "__tablename__": table_name,
            "__table_args__": (db.UniqueConstraint(*UNIQUE_COLUMNS, name=f"uq_{table_name}"),),
            "id": db.Column(db.Integer, primary_key=True),
            "tree_no": db.Column(db.String(50), nullable=False),
            "km_burji_no": db.Column(db.String(50), nullable=False),
            "kind_tree": db.Column(db.String(100), nullable=False),
            "girth": db.Column(db.Integer, nullable=False),
            "dia": db.Column(db.Integer),
            "cft": db.Column(db.Float),
            "shi_unit": db.Column(db.Float),
            "remarks": db.Column(db.String(255)),
            # where the person stood when the tree was entered (from the phone's GPS)
            "latitude": db.Column(db.Numeric(9, 6)),
            "longitude": db.Column(db.Numeric(10, 6)),
            "accuracy_m": db.Column(db.Float),  # GPS accuracy in metres (smaller is better)
            "to_dict": lambda self: {c.name: getattr(self, c.name) for c in self.__table__.columns},
        }
        model = type(f"Trees_{table_name}", (db.Model,), attrs)
        _tree_models[table_name] = model
    return model


def _slug(name):
    return re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")[:40]


def get_beat(name, owner_id):
    """The beat called `name` that belongs to owner_id (None = beats with no owner yet)."""
    return Beat.query.filter_by(name=name, user_id=owner_id).first()


LOCATION_COLUMNS = [("latitude", "DECIMAL(9,6)"), ("longitude", "DECIMAL(10,6)"), ("accuracy_m", "FLOAT")]
_location_checked = set()


def ensure_location_columns(table_name):
    """Older beat tables have no location columns: add them once (checked once per run)."""
    if table_name in _location_checked:
        return
    existing = {c["name"] for c in inspect(db.engine).get_columns(table_name)}
    for col, ddl in LOCATION_COLUMNS:
        if col in existing:
            continue
        try:
            with db.engine.begin() as conn:
                conn.execute(text(f"ALTER TABLE `{table_name}` ADD COLUMN `{col}` {ddl} NULL"))
            print(f"[database] added column {table_name}.{col}")
        except Exception:  # another worker may have added it a moment earlier
            if col not in {c["name"] for c in inspect(db.engine).get_columns(table_name)}:
                raise
    _location_checked.add(table_name)


def ensure_table(beat):
    try:
        tree_model(beat.table_name).__table__.create(db.engine, checkfirst=True)
    except Exception as exc:  # another worker may have created it a moment earlier
        if not inspect(db.engine).has_table(beat.table_name):
            raise exc
    ensure_location_columns(beat.table_name)


def get_or_create_beat(name, owner_id):
    """Find the owner's beat, or register it and create its own table."""
    beat = get_beat(name, owner_id)
    if beat is None:
        try:
            beat = Beat(name=name, user_id=owner_id, table_name="pending_" + uuid.uuid4().hex[:12])
            db.session.add(beat)
            db.session.flush()
            slug = _slug(name)
            beat.table_name = f"beat_{beat.id}" + (f"_{slug}" if slug else "")
            db.session.commit()
        except IntegrityError:
            db.session.rollback()
            beat = get_beat(name, owner_id)
    ensure_table(beat)
    return beat


def duplicate_message(tree_no, kind_tree, km_burji_no, beat):
    where = f"km / RD / burji no {km_burji_no} of {beat}"
    if TREE_NO_UNIQUE_PER_KIND:
        return (
            f"Tree no {tree_no} ({kind_tree.capitalize()}) is already saved in {where}. "
            "Use a different tree no."
        )
    return (
        f"Tree no {tree_no} is already saved in {where}. "
        "The same number can be used in another km, but not twice in the same km."
    )


def get_beat_options(owner_id):
    """Beats from the BEATS list plus every beat this user already has."""
    mine = {name for (name,) in db.session.query(Beat.name).filter(Beat.user_id == owner_id).all()}
    return sorted(set(BEATS) | mine, key=str.lower)


# Volume table as plain JSON, used by the page to preview dia / cft / shi unit live
VOLUME_JSON = {f"{kind}|{girth}": v for (kind, girth), v in VOLUME_TABLE.items()}


# ---------------------------------------------------------------------------
# Abstract: number of trees, total cft and total shi unit for each kind
# ---------------------------------------------------------------------------
def _new_abstract():
    return {"rows": {}, "trees": 0, "cft": 0.0, "shi_unit": 0.0}


def _add_to_abstract(a, kind, trees, cft, shi_unit):
    row = a["rows"].setdefault(kind, {"kind": kind, "trees": 0, "cft": 0.0, "shi_unit": 0.0})
    row["trees"] += trees
    row["cft"] += cft
    row["shi_unit"] += shi_unit
    a["trees"] += trees
    a["cft"] += cft
    a["shi_unit"] += shi_unit


def _finish_abstract(a):
    rows = sorted(a["rows"].values(), key=lambda r: r["kind"])
    for r in rows:
        r["cft"] = round(r["cft"], 3)
        r["shi_unit"] = round(r["shi_unit"], 3)
    return {
        "rows": rows,
        "trees": a["trees"],
        "cft": round(a["cft"], 3),
        "shi_unit": round(a["shi_unit"], 3),
    }


def get_abstracts(owner_id, only_beat=None):
    """Returns ({beat name: abstract}, grand_abstract) for ONE user. Empty beats are left out."""
    query = Beat.query.filter_by(user_id=owner_id)
    if only_beat:
        query = query.filter_by(name=only_beat)
    per_beat, grand = {}, _new_abstract()
    for beat in query.all():
        Tree = tree_model(beat.table_name)
        rows = db.session.query(
            Tree.kind_tree,
            func.count(Tree.id),
            func.coalesce(func.sum(Tree.cft), 0),
            func.coalesce(func.sum(Tree.shi_unit), 0),
        ).group_by(Tree.kind_tree).all()
        if not rows:
            continue
        a = _new_abstract()
        for kind, trees, cft, shi_unit in rows:
            _add_to_abstract(a, kind, trees, float(cft), float(shi_unit))
            _add_to_abstract(grand, kind, trees, float(cft), float(shi_unit))
        per_beat[beat.name] = _finish_abstract(a)
    return per_beat, _finish_abstract(grand)


# ---------------------------------------------------------------------------
# Accounts: sign in state, CSRF protection, register / login / logout
# ---------------------------------------------------------------------------
USERNAME_RE = re.compile(r"^[a-z0-9_.]{3,30}$")
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]{2,}$")
LOGIN_MAX_FAILS = 5          # wrong passwords allowed ...
LOGIN_WINDOW_SECONDS = 300   # ... within this many seconds, then wait
_failed_logins = defaultdict(list)
_DUMMY_HASH = generate_password_hash("not-a-real-password")  # same timing for unknown users


def csrf_token():
    """One secret token per session. Every POST form must send it back."""
    if "csrf_token" not in session:
        session["csrf_token"] = secrets.token_hex(16)
    return session["csrf_token"]


app.jinja_env.globals["csrf_token"] = csrf_token


@app.before_request
def load_user_and_check_csrf():
    g.user = None
    if request.endpoint in ("static", "database_health"):
        return

    user_id = session.get("user_id")
    if user_id:
        g.user = db.session.get(User, user_id)
        if g.user is None:  # the account no longer exists
            session.clear()

    if request.method == "POST":
        sent = request.form.get("csrf_token", "")
        expected = session.get("csrf_token", "")
        if not expected or not hmac.compare_digest(sent, expected):
            abort(400, "Your page has expired. Go back, refresh it and try again.")


@app.context_processor
def inject_user():
    return {"user": g.get("user")}


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if g.user is None:
            flash("Please sign in to continue.", "error")
            return redirect(url_for("login"))
        return view(*args, **kwargs)

    return wrapped


def too_many_failed_logins(key):
    now = time.time()
    recent = [t for t in _failed_logins[key] if now - t < LOGIN_WINDOW_SECONDS]
    _failed_logins[key] = recent
    return len(recent) >= LOGIN_MAX_FAILS


@app.route("/register", methods=["GET", "POST"])
def register():
    if g.user:
        return redirect(url_for("index"))

    error, form = None, {}
    if request.method == "POST":
        form = request.form
        full_name = " ".join(request.form.get("full_name", "").split())
        username = request.form.get("username", "").strip().lower()
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        confirm = request.form.get("confirm_password", "")

        if len(full_name) < 2 or len(full_name) > 100:
            error = "Please enter your full name."
        elif not USERNAME_RE.match(username):
            error = "Username must be 3 to 30 characters: small letters, numbers, dot or underscore."
        elif not EMAIL_RE.match(email) or len(email) > 120:
            error = "Please enter a valid email address."
        elif len(password) < 8:
            error = "Password must be at least 8 characters."
        elif len(password) > 128:
            error = "Password is too long (maximum 128 characters)."
        elif password.lower() == username:
            error = "Password must not be the same as the username."
        elif password != confirm:
            error = "The two passwords do not match."
        elif User.query.filter_by(username=username).first():
            error = "That username is already taken. Please choose another."
        elif User.query.filter_by(email=email).first():
            error = "That email is already registered. Try signing in instead."
        else:
            is_first_user = User.query.count() == 0
            user = User(
                full_name=full_name,
                username=username,
                email=email,
                password_hash=generate_password_hash(password),
            )
            db.session.add(user)
            try:
                db.session.flush()
                claimed = 0
                if is_first_user:
                    # Beats (and their trees) saved before accounts existed go to the first account
                    claimed = (
                        Beat.query.filter(Beat.user_id.is_(None))
                        .update({"user_id": user.id}, synchronize_session=False)
                    )
                db.session.commit()
            except IntegrityError:
                db.session.rollback()
                error = "That username or email is already registered."
            else:
                session.clear()
                session["user_id"] = user.id
                message = f"Welcome, {user.full_name}. Your account is ready."
                if claimed:
                    message += f" The {claimed} existing beats were added to your account."
                flash(message, "success")
                return redirect(url_for("index"))

    return render_template("register.html", error=error, form=form)


@app.route("/login", methods=["GET", "POST"])
def login():
    if g.user:
        return redirect(url_for("index"))

    error, form = None, {}
    if request.method == "POST":
        form = request.form
        identifier = request.form.get("username", "").strip().lower()  # username or email
        password = request.form.get("password", "")
        key = (request.remote_addr, identifier)

        if too_many_failed_logins(key):
            error = "Too many wrong attempts. Please wait 5 minutes and try again."
        else:
            if "@" in identifier:
                user = User.query.filter_by(email=identifier).first()
            else:
                user = User.query.filter_by(username=identifier).first()
            password_ok = check_password_hash(user.password_hash if user else _DUMMY_HASH, password)
            if user and password_ok:
                _failed_logins.pop(key, None)
                session.clear()  # new session after sign in
                session["user_id"] = user.id
                return redirect(url_for("index"))
            _failed_logins[key].append(time.time())
            error = "Wrong username / email or password."

    return render_template("login.html", error=error, form=form)


@app.route("/logout", methods=["POST"])
def logout():
    session.clear()
    flash("You have signed out.", "success")
    return redirect(url_for("login"))


# ---------------------------------------------------------------------------
# Helpers for the tree form
# ---------------------------------------------------------------------------
def normalize_beat(name, owner_id):
    """Trim spaces and reuse the spelling of an existing beat (case-insensitive)."""
    beat = " ".join((name or "").split())
    for known in get_beat_options(owner_id):
        if known.lower() == beat.lower():
            return known
    return beat


def read_tree_form(form, beat, exclude_id=None):
    """Check the form. Returns (error, data); data is ready to save when error is None.

    beat is the Beat whose table the tree goes into (or None if no beat is chosen).
    exclude_id is the tree being edited, so it is not counted as its own duplicate.
    """
    tree_no = " ".join(form.get("tree_no", "").split())
    km_burji_no = " ".join(form.get("km_burji_no", "").split())
    kind_tree = form.get("kind_tree", "").strip().lower()
    girth_raw = form.get("girth", "").strip()
    remarks = form.get("remarks", "").strip()

    if beat is None:
        return "Please set the beat name first.", None
    if not (tree_no and km_burji_no and kind_tree and girth_raw):
        return "Please enter the tree no, km / RD / burji no, select a tree kind and enter the girth.", None

    try:
        girth = int(girth_raw)
    except ValueError:
        return "Girth must be a whole number.", None

    # Location from the phone's GPS (optional): both numbers or neither
    latitude = longitude = accuracy = None
    lat_raw = form.get("latitude", "").strip()
    lon_raw = form.get("longitude", "").strip()
    acc_raw = form.get("accuracy_m", "").strip()
    if lat_raw or lon_raw:
        try:
            lat, lon = float(lat_raw), float(lon_raw)
        except ValueError:
            return "The location is not valid. Capture it again or clear it.", None
        if not (-90 <= lat <= 90 and -180 <= lon <= 180):
            return "The location is not valid. Capture it again or clear it.", None
        latitude, longitude = round(lat, 6), round(lon, 6)
        try:
            acc = float(acc_raw)
            accuracy = round(acc, 1) if 0 <= acc < 100000 else None
        except ValueError:
            accuracy = None
    elif REQUIRE_LOCATION:
        return "Please capture the location before saving.", None

    row = VOLUME_TABLE.get((kind_tree, girth))
    if row is None:
        return f"No data found for {kind_tree} with girth {girth}.", None

    Tree = tree_model(beat.table_name)
    key = {"km_burji_no": km_burji_no, "kind_tree": kind_tree, "tree_no": tree_no}
    duplicate = Tree.query.filter_by(**{col: key[col] for col in UNIQUE_COLUMNS})
    if exclude_id is not None:
        duplicate = duplicate.filter(Tree.id != exclude_id)
    if duplicate.first():
        return duplicate_message(tree_no, kind_tree, km_burji_no, beat.name), None

    return None, {
        "tree_no": tree_no,
        "km_burji_no": km_burji_no,
        "kind_tree": kind_tree,
        "girth": girth,
        "dia": row["dia"],
        "cft": row["cft"],
        "shi_unit": row["shi_unit"],
        "remarks": remarks or None,
        "latitude": latitude,
        "longitude": longitude,
        "accuracy_m": accuracy,
    }


def my_beat_or_404(beat_id):
    """A beat of the signed-in user. Someone else's beat looks like it does not exist."""
    return Beat.query.filter_by(id=beat_id, user_id=g.user.id).first_or_404()


def render_index(error=None, form=None, edit_tree=None, edit_beat=None):
    """Build the page: one section per beat of this user (current beat first, then A-Z)."""
    current_beat = session.get("beat_name")

    abstracts, grand = get_abstracts(g.user.id)
    sections = []
    for beat in Beat.query.filter_by(user_id=g.user.id).all():
        if beat.name not in abstracts:
            continue
        Tree = tree_model(beat.table_name)
        sections.append({
            "id": beat.id,
            "beat": beat.name,
            "trees": Tree.query.order_by(Tree.id.asc()).all(),
            "abstract": abstracts[beat.name],
        })
    sections.sort(key=lambda s: (s["beat"] != current_beat, s["beat"].lower()))

    return render_template(
        "index.html",
        kinds=KINDS,
        beats=get_beat_options(g.user.id),
        current_beat=current_beat,
        sections=sections,
        grand=grand,
        volume=VOLUME_JSON,
        error=error,
        form=form or {},
        edit_tree=edit_tree,
        edit_beat=edit_beat,
    )


# ---------------------------------------------------------------------------
# Routes
# ---------------------------------------------------------------------------
@app.route("/set-beat", methods=["POST"])
@login_required
def set_beat():
    """Choose the beat once. Its own table is created the first time it is used."""
    name = normalize_beat(request.form.get("beat_name"), g.user.id)

    if not name:
        flash("Please pick or type a beat name.", "error")
        return redirect(url_for("index"))
    if len(name) > 100:
        flash("Beat name is too long (maximum 100 characters).", "error")
        return redirect(url_for("index"))

    get_or_create_beat(name, g.user.id)
    session["beat_name"] = name
    return redirect(url_for("index"))


@app.route("/", methods=["GET", "POST"])
@login_required
def index():
    current_beat = session.get("beat_name")

    # ----- Add a new tree -----
    if request.method == "POST":
        beat = get_or_create_beat(current_beat, g.user.id) if current_beat else None
        error, data = read_tree_form(request.form, beat)

        if error is None:
            Tree = tree_model(beat.table_name)
            tree = Tree(**data)
            db.session.add(tree)
            try:
                db.session.commit()
            except IntegrityError:
                db.session.rollback()
                error = duplicate_message(
                    data["tree_no"], data["kind_tree"], data["km_burji_no"], beat.name
                )
            else:
                flash(
                    f"Saved tree no {tree.tree_no} ({tree.km_burji_no}) in {beat.name}: "
                    f"{tree.kind_tree.capitalize()}, girth {tree.girth}, dia {tree.dia}, "
                    f"cft {tree.cft}, shi unit {tree.shi_unit}",
                    "success",
                )
                # Redirect so refreshing the page does not save a duplicate
                return redirect(url_for("index"))

        return render_index(error=error, form=request.form)

    # ----- Show the page, optionally with one tree loaded for editing (?edit=ID&beat=BEAT_ID) -----
    edit_tree = edit_beat = None
    edit_id = request.args.get("edit", type=int)
    beat_id = request.args.get("beat", type=int)
    if edit_id and beat_id:
        edit_beat = Beat.query.filter_by(id=beat_id, user_id=g.user.id).first()
        if edit_beat:
            edit_tree = db.session.get(tree_model(edit_beat.table_name), edit_id)
        if edit_tree is None:
            flash("That tree no longer exists.", "error")
            return redirect(url_for("index"))

    return render_index(
        form=edit_tree.to_dict() if edit_tree else {}, edit_tree=edit_tree, edit_beat=edit_beat
    )


@app.route("/update/<int:beat_id>/<int:tree_id>", methods=["POST"])
@login_required
def update_tree(beat_id, tree_id):
    """Save changes to a tree in its beat's table. dia, cft and shi unit are worked out again."""
    beat = my_beat_or_404(beat_id)
    Tree = tree_model(beat.table_name)
    tree = db.get_or_404(Tree, tree_id)

    error, data = read_tree_form(request.form, beat, exclude_id=tree.id)

    if error is None:
        for field, value in data.items():
            setattr(tree, field, value)
        try:
            db.session.commit()
        except IntegrityError:
            db.session.rollback()
            error = duplicate_message(
                data["tree_no"], data["kind_tree"], data["km_burji_no"], beat.name
            )
        else:
            flash(
                f"Updated tree no {tree.tree_no} ({tree.km_burji_no}) in {beat.name}: "
                f"{tree.kind_tree.capitalize()}, girth {tree.girth}, dia {tree.dia}, "
                f"cft {tree.cft}, shi unit {tree.shi_unit}",
                "success",
            )
            return redirect(url_for("index"))

    db.session.rollback()
    return render_index(error=error, form=request.form, edit_tree=tree, edit_beat=beat)


@app.route("/delete/<int:beat_id>/<int:tree_id>", methods=["POST"])
@login_required
def delete_tree(beat_id, tree_id):
    beat = my_beat_or_404(beat_id)
    tree = db.get_or_404(tree_model(beat.table_name), tree_id)
    label = (
        f"tree no {tree.tree_no} ({tree.kind_tree.capitalize()}, km / burji no "
        f"{tree.km_burji_no}) from {beat.name}"
    )
    db.session.delete(tree)
    db.session.commit()
    flash(f"Deleted {label}.", "success")
    return redirect(url_for("index"))


# ---------------------------------------------------------------------------
# Excel download: one separate sheet per beat, abstract at the bottom of each
# The beat name is written once at the top of its sheet.
# ---------------------------------------------------------------------------
EXCEL_HEADERS = [
    "ID", "Tree No", "Km / RD / Burji No", "Kind",
    "Girth", "Dia", "CFT", "Shi Unit", "Latitude", "Longitude", "Accuracy (m)", "Remarks",
]
EXCEL_WIDTHS = [14, 14, 16, 14, 10, 8, 10, 12, 12, 12, 14, 40]


def _num(value):
    return float(value) if value is not None else None
TOTAL_BORDER = Border(top=Side(style="double"))


def safe_sheet_name(name, used):
    """Excel sheet names: max 31 chars, no []:*?/\\ and must be unique."""
    base = re.sub(r"[\[\]:*?/\\]", "", name).strip()[:31] or "Beat"
    candidate, n = base, 2
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def _bold_last_row(ws, border=None):
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
        if border:
            cell.border = border


def write_abstract_block(ws, abstract, title="Abstract"):
    ws.append([])
    ws.append([title])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
    ws.append(["No. of trees", "Kind", "CFT", "Shi unit"])
    _bold_last_row(ws)
    for r in abstract["rows"]:
        ws.append([r["trees"], r["kind"].capitalize(), r["cft"], r["shi_unit"]])
    ws.append([abstract["trees"], "Total", abstract["cft"], abstract["shi_unit"]])
    _bold_last_row(ws, TOTAL_BORDER)


def write_beat_sheet(ws, beat, trees, abstract):
    """Row 1: beat name. Row 2: column names. Then the trees, then the abstract."""
    ws.append([f"Beat: {beat}"])
    ws["A1"].font = Font(bold=True, size=14)

    ws.append(EXCEL_HEADERS)
    _bold_last_row(ws)

    for t in trees:
        ws.append([
            t.id, t.tree_no, t.km_burji_no, t.kind_tree.capitalize(),
            t.girth, t.dia, t.cft, t.shi_unit,
            _num(t.latitude), _num(t.longitude), t.accuracy_m, t.remarks,
        ])

    write_abstract_block(ws, abstract)

    for col, width in zip("ABCDEFGHIJKL", EXCEL_WIDTHS):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A3"


def write_grand_sheet(ws, grand, per_beat):
    ws.append(["Grand abstract: all beats"])
    ws["A1"].font = Font(bold=True, size=14)
    write_abstract_block(ws, grand, title="By kind")

    ws.append([])
    ws.append(["By beat"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True, size=12)
    ws.append(["No. of trees", "Beat", "CFT", "Shi unit"])
    _bold_last_row(ws)
    for beat in sorted(per_beat, key=str.lower):
        a = per_beat[beat]
        ws.append([a["trees"], beat, a["cft"], a["shi_unit"]])
    ws.append([grand["trees"], "Total", grand["cft"], grand["shi_unit"]])
    _bold_last_row(ws, TOTAL_BORDER)

    for col, width in zip("ABCD", [14, 24, 12, 12]):
        ws.column_dimensions[col].width = width


@app.route("/download-excel")
@login_required
def download_excel():
    """All of this user's beats (one sheet each). /download-excel?beat=Name for a single beat."""
    only_beat = request.args.get("beat")
    per_beat, grand = get_abstracts(g.user.id, only_beat)

    wb = Workbook()
    first_sheet = wb.active

    beats = sorted(per_beat, key=str.lower)
    if not beats:
        first_sheet.title = "No Data"
    used = set()
    for i, name in enumerate(beats):
        Tree = tree_model(get_beat(name, g.user.id).table_name)
        trees = Tree.query.order_by(Tree.id).all()
        ws = first_sheet if i == 0 else wb.create_sheet()
        ws.title = safe_sheet_name(name, used)
        write_beat_sheet(ws, name, trees, per_beat[name])

    if len(beats) > 1:
        ws = wb.create_sheet()
        ws.title = safe_sheet_name("Grand Abstract", used)
        write_grand_sheet(ws, grand, per_beat)

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    label = re.sub(r"[^\w\-]+", "_", only_beat) if only_beat else "all_beats"
    filename = f"tree_data_{label}_{datetime.now():%Y-%m-%d_%H%M}.xlsx"
    return send_file(
        buffer,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ---------------------------------------------------------------------------
# Database repair, runs at start and is safe to run again and again
#  - "beats" gets its owner column (user_id)
#  - the old rule "beat name must be unique for everybody" is replaced by
#    "beat name must be unique for each user"
# ---------------------------------------------------------------------------
def upgrade_database():
    if db.engine.dialect.name != "mysql":
        return  # a fresh local SQLite file is created with the right layout already

    insp = inspect(db.engine)
    if not insp.has_table("beats"):
        return

    if "user_id" not in {c["name"] for c in insp.get_columns("beats")}:
        with db.engine.begin() as conn:
            conn.execute(text("ALTER TABLE beats ADD COLUMN user_id INT NULL"))
        print("[database] added column beats.user_id")

    indexes = inspect(db.engine).get_indexes("beats")
    for idx in indexes:  # drop the old "name must be unique" rule
        if idx.get("unique") and list(idx["column_names"]) == ["name"]:
            with db.engine.begin() as conn:
                conn.execute(text(f"ALTER TABLE beats DROP INDEX `{idx['name']}`"))
            print(f"[database] removed old unique rule `{idx['name']}` on beats.name")

    has_new_rule = any(
        idx.get("unique") and list(idx["column_names"]) == ["user_id", "name"] for idx in indexes
    )
    if not has_new_rule:
        with db.engine.begin() as conn:
            conn.execute(text("ALTER TABLE beats ADD CONSTRAINT uq_beats_user_name UNIQUE (user_id, name)"))
        print("[database] added unique rule on beats (user_id, name)")

    fk_columns = [c for fk in inspect(db.engine).get_foreign_keys("beats") for c in fk["constrained_columns"]]
    if "user_id" not in fk_columns:
        try:
            with db.engine.begin() as conn:
                conn.execute(text(
                    "ALTER TABLE beats ADD CONSTRAINT fk_beats_user FOREIGN KEY (user_id) REFERENCES app_user (id)"
                ))
            print("[database] linked beats to users")
        except Exception as exc:
            print(f"[database] note: could not add the user link ({exc.__class__.__name__}); "
                  "the app still keeps each user's data separate.")


# ---------------------------------------------------------------------------
# Move trees from the OLD single table ("shisham" or "trees") into the per-beat
# tables. Runs at start. The old table is renamed to <name>_old_backup, not deleted.
# The moved beats have no owner until the first account is registered.
# ---------------------------------------------------------------------------
def migrate_old_tables():
    insp = inspect(db.engine)
    for old in ("shisham", "trees"):
        if not insp.has_table(old):
            continue
        cols = {c["name"] for c in insp.get_columns(old)}
        if not {"tree_no", "kind_tree", "girth"} <= cols:
            continue
        rows = db.session.execute(text(f"SELECT * FROM `{old}`")).mappings().all()
        moved = skipped = 0
        for r in rows:
            name = " ".join(((r.get("beat_name") if "beat_name" in cols else None) or "Unassigned").split())
            beat = get_or_create_beat(normalize_beat(name, None), None)
            Tree = tree_model(beat.table_name)
            db.session.add(Tree(
                tree_no=r["tree_no"], km_burji_no=r.get("km_burji_no") or "-",
                kind_tree=r["kind_tree"], girth=r["girth"], dia=r.get("dia"),
                cft=r.get("cft"), shi_unit=r.get("shi_unit"), remarks=r.get("remarks"),
            ))
            try:
                db.session.commit()
                moved += 1
            except IntegrityError:
                db.session.rollback()
                skipped += 1
        backup = f"{old}_old_backup"
        if not insp.has_table(backup):
            with db.engine.begin() as conn:
                conn.execute(text(f"ALTER TABLE `{old}` RENAME TO `{backup}`"))
        print(f"[database] moved {moved} trees from `{old}` into per-beat tables "
              f"({skipped} skipped as duplicates). Old table kept as `{backup}`.")


# ---------------------------------------------------------------------------
# Friendly page when the database cannot be reached
# Error details are only shown while SHOW_ERRORS=1 is set (turn it on to debug,
# then remove it: details can contain server names).
# ---------------------------------------------------------------------------
SHOW_ERRORS = os.getenv("SHOW_ERRORS") == "1"

DB_ERROR_PAGE = """<!doctype html><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Database not reachable</title>
<body style="font:16px/1.5 system-ui,sans-serif;max-width:640px;margin:40px auto;padding:0 16px">
<h1>Cannot connect to the database</h1>
<p>The app is running, but it could not reach MySQL. Check these:</p>
<ol>
<li><b>Host</b>: copy it exactly from Aiven (it starts with <code>mysql-</code>).</li>
<li><b>Port, user, password, database name</b> in your <code>.env</code> file.</li>
<li><b>Aiven service</b> is running and your IP is in <i>Allowed IP addresses</i>.</li>
<li><b>ca.pem</b> is next to <code>app.py</code>.</li>
</ol>
<p>Open <a href="/health/db">/health/db</a> to see whether the database answers.</p>
{% if detail %}<pre style="background:#f3f3f3;padding:10px;overflow:auto">{{ detail }}</pre>{% endif %}</body>"""


@app.errorhandler(OperationalError)
def database_unreachable(exc):
    db.session.rollback()
    app.logger.error("Database error: %s", exc)
    detail = (str(exc.orig) if getattr(exc, "orig", None) else str(exc))[:300] if SHOW_ERRORS else ""
    return render_template_string(DB_ERROR_PAGE, detail=detail), 503


@app.errorhandler(Exception)
def show_error(exc):
    if isinstance(exc, HTTPException):
        return exc
    db.session.rollback()
    app.logger.exception("Unhandled error")
    if SHOW_ERRORS:
        return f"<h1>Error</h1><pre>{escape(exc.__class__.__name__)}: {escape(str(exc)[:500])}</pre>", 500
    return "Internal Server Error", 500


@app.route("/health/db")
def database_health():
    try:
        with db.engine.connect() as connection:
            connection.execute(text("SELECT 1"))
        return {"database": "connected", "engine": db.engine.dialect.name}, 200
    except Exception as exc:
        app.logger.exception("Database connection failed")
        result = {"database": "connection failed", "error": exc.__class__.__name__}
        if SHOW_ERRORS:
            result["detail"] = str(getattr(exc, "orig", exc))[:300]
        return result, 500


# ---------------------------------------------------------------------------
# Create / repair tables on start (also works under gunicorn)
# ---------------------------------------------------------------------------
def ensure_all_beat_tables():
    for beat in Beat.query.all():
        ensure_table(beat)


def init_database():
    with app.app_context():
        steps = [
            ("create tables", db.create_all),          # app_user, beats
            ("repair beats table", upgrade_database),  # owner column + new unique rule
            ("create beat tables", ensure_all_beat_tables),
            ("move old data", migrate_old_tables),
        ]
        for label, step in steps:
            try:
                step()
            except Exception as exc:
                db.session.rollback()
                print(f"[database] could not {label}: {exc.__class__.__name__}: "
                      f"{str(exc).splitlines()[0][:200]}")


init_database()


if __name__ == "__main__":
    app.run(debug=os.getenv("FLASK_DEBUG", "0") == "1", port=int(os.getenv("PORT", "5000")))
